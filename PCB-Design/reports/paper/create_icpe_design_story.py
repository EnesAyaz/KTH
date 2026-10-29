from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import re

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'icpe'/'SPB_ICPE_design_story_draft_2026-10-23.docx'
AS=ROOT/'icpe'/'design_story_assets'; AS.mkdir(exist_ok=True)
if OUT.exists(): raise FileExistsError(OUT)
doc=Document()
s=doc.sections[0]; s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=doc.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.03
doc.styles['Title'].font.size=Pt(18)
doc.styles['Heading 1'].font.size=Pt(13)
doc.styles['Heading 2'].font.size=Pt(11)
doc.styles['Caption'].font.size=Pt(9)
def p(t): return doc.add_paragraph(t)
def h(t): doc.add_heading(t,1)
def page(t): doc.add_page_break(); h(t)
def fig(path,cap,width=6.5):
 pa=doc.add_paragraph(); pa.alignment=1
 pa.add_run().add_picture(str(path),width=Inches(width)); pa.paragraph_format.space_after=Pt(2)
 doc.add_paragraph(cap,'Caption')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
 for c,txt in zip(t.rows[0].cells,headers):
  c.text=txt
  sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7EDF3'); c._tc.get_or_add_tcPr().append(sh)
 for row in rows:
  for c,txt in zip(t.add_row().cells,row): c.text=str(txt)
 for row in t.rows:
  for c in row.cells:
   for pa in c.paragraphs:
    pa.paragraph_format.space_after=Pt(4); pa.paragraph_format.space_before=Pt(4)
    for r in pa.runs: r.font.size=Pt(9)
  trPr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trPr.append(no)
 return t

im=Image.new('RGB',(1600,950),'white'); dr=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',27)
small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',23)
boxes=[(80,40,680,150,'Traction specification\n1200 V  300 kW  50 kHz'),(80,215,680,325,'Compare voltage classes and D FoM\nSelect 100 V EPC2361'),(80,390,680,500,'Sweep parallel count and loop inductance\nLoss  temperature  voltage margin'),(80,565,680,675,'Four devices per switch\nSymmetric pairs and Kelvin returns'),(920,40,1520,150,'KiCad layout\nPower cells and gate drive'),(920,215,1520,325,'Ansys Q3D extraction\nPower and gate R L matrices'),(920,390,1520,500,'LTspice double pulse simulation\nCurrent sharing  losses  VDS  VGS'),(920,650,1520,760,'Check design targets\nRevise layout and drive if needed'),(920,825,1520,935,'DPT board and experimental correlation\nPlanned for the full paper')]
for x1,y1,x2,y2,txt in boxes:
 dr.rounded_rectangle((x1,y1,x2,y2),radius=12,fill='#EDF2F7',outline='#263D50',width=3)
 dr.multiline_text(((x1+x2)/2,(y1+y2)/2),txt,font=font,fill='black',anchor='mm',align='center',spacing=10)
def arrow(a,b):
 dr.line((a,b),fill='#263D50',width=5)
 x,y=b; dx=x-a[0];dy=y-a[1]
 if abs(dy)>abs(dx): dr.polygon([(x,y),(x-10,y-18 if dy>0 else y+18),(x+10,y-18 if dy>0 else y+18)],fill='#263D50')
 else: dr.polygon([(x,y),(x-18 if dx>0 else x+18,y-10),(x-18 if dx>0 else x+18,y+10)],fill='#263D50')
for x in [380]:
 for a,b in [(150,215),(325,390),(500,565)]: arrow((x,a),(x,b))
arrow((680,620),(800,620)); dr.line((800,620,800,95),fill='#263D50',width=5);arrow((800,95),(920,95))
for a,b in [(150,215),(325,390),(500,650),(760,825)]: arrow((1220,a),(1220,b))
dr.line((1520,705,1570,705,1570,95),fill='#263D50',width=4);arrow((1570,95),(1520,95))
dr.text((1270,565),'Iterate until constraints are met',font=small,fill='black',anchor='mm')
im.save(AS/'design_flow.png')

doc.add_paragraph('Building Block Design for Stacked Polyphase Bridge Converters for High Voltage Traction Drives','Title')
p('Enes Ayaz\nDivision of Electric Power and Energy Systems\nKTH Royal Institute of Technology  Stockholm  Sweden').alignment=1
h('Abstract')
p('A stacked polyphase bridge (SPB) distributes a high traction dc voltage across series connected three phase submodules, enabling low voltage devices while retaining the required system power. This paper develops the phase leg building block of a 1200 V, 300 kW SPB through a coupled device, gate drive, and layout design procedure. A device figure of merit comparison motivates the 100 V EPC2361. Loss, thermal, and overshoot studies then select four parallel devices per switch at a nominal submodule voltage of 75 V. Dynamic current sharing motivates symmetric commutation cells, individual Kelvin returns, and one driver channel for each pair of parallel transistors. An iterative KiCad, Ansys Q3D, and LTspice workflow evaluates the power and gate networks together. Existing extraction results give 0.41 nH per cell for the operating board and 0.65 nH for the double pulse test board. Simulation identifies the additional gate loop inductance and measurement shunts as important constraints: adjusted turn-on resistance reduces the operating-board peak to 86 V, while the DPT configuration reaches 93 V and requires further margin improvement against the 90 V target. Experimental double pulse validation is planned for the full paper.')
p('Keywords: stacked polyphase bridge, traction inverter, GaN, parallel devices, gate drive, parasitic extraction, double pulse test.')
h('1 Introduction')
p('High voltage traction drives require an inverter architecture that combines manageable semiconductor stress with low loss and practical integration. In a conventional two level inverter, the devices must withstand the entire dc link and commutation overshoot. The SPB instead connects three phase bridges in series on the dc side, with each bridge supplying an electrically isolated winding set of a multiphase machine [1–3]. The architecture makes the submodule voltage, and therefore the semiconductor voltage class, a design choice. It can reduce the voltage step applied to each winding set and offers a modular route to integrated drives. These benefits depend on the machine, modulation, voltage balancing, insulation, and thermal design; they are not automatic consequences of stacking.')
p('Prior SPB work establishes the topology and modulation principles [1–3], dc-link stability [4], and system-level comparisons with alternative GaN converters [5]. The remaining hardware question addressed here is how to realize a repeatable low voltage, high current building block. Increasing the number of submodules reduces blocking voltage but does not reduce the phase current required at fixed system power. Low voltage GaN devices must consequently be paralleled, making commutation inductance, gate impedance, and dynamic sharing central design constraints [6,7].')
p('The contribution is a linked design procedure: select the device voltage class using a device figure of merit; select the parallel count using loss and voltage margin; translate sharing sensitivity into a symmetric gate and power layout; and close the design loop using extracted parasitics in switching simulations. The resulting DPT board provides the platform for experimental correlation in the full paper.')

page('2 From traction requirements to the device voltage class')
fig(ROOT/'icpe/img/fig_spb.png','Fig. 1. SPB architecture and phase-leg building block from the existing design study.',6.3)
p('The design point is 1200 V dc, 300 kW, 50 kHz switching frequency, modulation index M = 1.15, and power factor cos φ = 0.9. With balanced power among N submodules, Vm = Vdc/N and P = 3 M Vdc Iph,rms cos φ /(2√2). The resulting phase current is approximately 228 A rms and 322 A peak, independent of N. At N = 16 the nominal submodule voltage is 75 V. This is a candidate operating point to be checked against transient stress, rather than an independently established safe voltage.')
p('For screening, define the device figure of merit D-FoM = RDS(on) Coss,tr, with Coss,tr = Qoss/Vtest. Smaller values are preferable. This charge-equivalent capacitance is used consistently in the existing survey; it must not be confused with a small-signal capacitance at a single bias. Under a simplified conduction-plus-output-charge loss model, optimizing the continuous parallel count gives a minimum system loss proportional to Vdc Iph,rms √(fsw D-FoM). Voltage class cancels to first order, so the comparison should use actual candidate devices rather than assuming a higher or lower class is inherently superior.')
table(['Voltage class','Representative device','RDS typ mΩ','D-FoM ps','Nominal Vm / N'],[
['100 V','EPC2361','0.75','1.35','75 V / 16'],['200 V','EPC2215','6.0','4.14','150 V / 8'],['350 V','EPC2050','55','6.88','240 V / 5'],['650 V','IGT65R140D2','140','4.90','400 V / 3']])
doc.add_paragraph('Table I. Representative candidates from the existing ten-device survey. Qoss test voltages differ; D-FoM is an initial screening metric, not an operating-point loss prediction.','Caption')
p('The surveyed EPC2361 has the smallest D-FoM, together with a Kelvin source connection and a 3 × 5 mm package. Its manufacturer specifies 100 V blocking and 0.75 mΩ typical on-resistance [8]. It is therefore selected for detailed simulation. The benefit of sixteen low voltage submodules must ultimately be weighed against the added drivers, auxiliary supplies, interconnections, and machine winding sets. Nonlinear output charge, temperature dependence, dead-time loss, and realizable integer device counts are included in subsequent analysis rather than inferred from D-FoM alone.')

page('3 Selecting four parallel transistors per switch')
p('The next decision is the number n of transistors per high-side or low-side switch. At the peak phase current, n = 2, 3, and 4 imply approximately 161, 107, and 80.5 A per transistor before dynamic imbalance. Conduction loss falls with n, whereas switched capacitance, gate-drive demand, component count, and routing complexity increase. A useful screening interval for the commutation inductance is 0.5–1 nH per cell, spanning the intended compact layout and a less favorable measurement configuration.')
fig(ROOT/'icpe/img/flow_npar2.png','Fig. 2. Existing parallel-count study at 228 A rms and 50 kHz: semiconductor loss, estimated junction temperature, and peak drain voltage for assumed cell-loop inductances.',6.5)
table(['Devices per switch','Leg loss W at 0.5 nH','Estimated Tj °C','Peak VDS V at 0.5 / 1 nH'],[
['2','39.5','141','131 / 180'],['3','23.2','84','96 / 123'],['4','18.0','69','83 / 93'],['6','14.8','59','86 / 84'],['8','14.6','56','83 / 87']])
doc.add_paragraph('Table II. Existing lumped-model sweep. Temperature estimates use 40 °C ambient and the assumed load-board cooling path. Values are simulated or calculated, not measured.','Caption')
p('Four devices are selected because they meet the preliminary thermal objective and provide useful voltage margin in the 0.5 nH case, while allowing two identical gate-drive pairs. Three devices exceed the 90 V design target even at the smaller assumed inductance. Increasing from four to six devices saves about 3.2 W per leg in this model but adds components and drive complexity; four is a practical compromise rather than a global loss optimum. The nonmonotonic voltage entries show why the full switching waveform must be checked at each operating point.')
p('The first-order relation VDS,pk ≈ Vm + Lloop |di/dt| is used as a margin screen. At Vm = 75 V, a 90 V target allows 30 A/ns for 0.5 nH but only 15 A/ns for 1 nH. The inductance and current slope must refer to the same commutation path; a per-cell loop must not be combined with the aggregate four-cell current slope. The 75 V choice is retained conditionally, with the extracted-network simulation determining the final gate setting and remaining margin.')

page('4 Symmetric current paths and gate drive in pairs')
p('Parallel count alone does not ensure usable current capability. During the switching interval, differences in gate-path impedance, common-source inductance, threshold voltage, and driver timing can concentrate current and energy in one device. The existing two-device sensitivity study at 75 V and 120 A isolates these mechanisms [6,7]. A positive temperature coefficient of on-resistance helps static sharing, but does not eliminate nanosecond-scale dynamic imbalance.')
fig(ROOT/'figures/fig_sharing.png','Fig. 3. Simulated sensitivity of two-device current sharing to gate-loop mismatch and loss of the Kelvin source connection.',6.4)
table(['Perturbation in the existing study','Switching-energy imbalance','Design response'],[
['Power-path inductance +0.3 nH','2%','Mirror the commutation cells'],['Gate loops 2 versus 8 nH','19%','Equal gate branches and local returns'],['No Kelvin return and +0.2 nH common-source mismatch','29%','Separate Kelvin connection per device'],['Threshold offset +0.3 V','7%','Check device spread and driver timing']])
doc.add_paragraph('Table III. One perturbation at a time. The nominal symmetric result does not establish tolerance robustness.','Caption')
p('The selected arrangement uses two identical driver blocks per phase leg. Each 2EDF7275K has two output channels: one drives two high-side transistors, and the other drives the corresponding two low-side transistors. Thus one channel serves two parallel transistors; two dual-channel driver ICs serve the eight transistors in the complete leg. The 4 A source and 8 A sink capability [9] supports rapid gate-charge removal and suppresses Miller-induced gate rise. The existing four-device-per-channel study indicates approximately 31% more turn-on energy and a peak sink demand above 8 A, motivating division into pairs.')
p('Each pair is arranged with point symmetry around its driver. Star branches, individual gate resistors, and individual Kelvin returns aim to equalize impedance without routing power current through the gate reference. Strong sink capability must be assessed together with common gate inductance and output resistance: the driver current rating alone does not guarantee low device-terminal impedance. Timing mismatch between the two driver blocks and semiconductor parameter spread remain part of the required robustness checks.')

page('5 Iterative layout extraction and switching simulation')
fig(AS/'design_flow.png','Fig. 4. Proposed coupled design workflow. Failure of a voltage, gate, sharing, or thermal constraint returns the design to layout or gate-drive adjustment.',6.4)
p('The KiCad implementation places four local commutation cells as two symmetric pairs. Each cell contains one high-side and one low-side transistor with local bypass capacitors. Adjacent power planes separated by a thin dielectric shorten the commutation path, while gate and Kelvin tracks run together on adjacent layers. The existing board uses four copper layers with 2 oz copper and approximately 0.1 mm separation between the top power layers. Local bypassing closes the fast switching path; the larger dc-link bank supplies the lower-frequency switching ripple.')
p('The routed copper is transferred to Ansys Q3D to extract resistance and partial-inductance matrices, including mutual terms. The existing procedure uses the cell region at 100 MHz for high-frequency loop extraction and the full board at dc for resistive paths. The gate extraction includes the common driver-to-star path and individual branches. Capacitor ESL, shunt inductance, and component interconnects are added consistently to avoid omitting or double-counting parasitics. Per-cell loop inductance is evaluated under the simultaneous four-cell commutation condition.')
p('LTspice combines the extracted network with eight EPC2361 models, two driver blocks, capacitor bias derating, and the appropriate test shunts. Each iteration evaluates individual device currents, switching energies, both drain voltages, and gate-to-Kelvin-source voltages. Layout revisions target both absolute inductance and matching. Driver resistance, turn-off impedance, and off-state bias are then adjusted within the device and driver limits. Frequency-dependent extraction and model bandwidth should be checked against the observed ringing spectrum before experimental correlation.')
p('The acceptance targets are a nominal peak VDS no greater than 90 V, gate excursions within the device limits, no simulated false turn-on, acceptable device energy spread, and thermal compatibility with the assumed cooling path. A symmetric result in the nominal model is only the starting point for sweeps of temperature, device parameters, driver skew, and parasitic tolerances.')

page('6 Loss distribution and parasitic limits')
table(['Quantity','Operating board','DPT board'],[
['Effective commutation loop per cell','0.41 nH','0.65 nH'],['Spread among cell loops','1.6%','0.8%'],['Gate loop for simultaneous pair switching','Not established here','20.6 nH high side / 12.3 nH low side'],['Dominant added measurement element','No cell shunt','Coaxial shunt and connection']])
doc.add_paragraph('Table IV. Existing Q3D results. Effective pair gate-loop values include the shared path and must not be read as individual branch inductances.','Caption')
p('The extracted operating-board loop comprises approximately 0.35 nH of copper and 0.06 nH of bypass-capacitor ESL. On the DPT board, copper, ESL, and the shunt contribute approximately 0.47, 0.06, and 0.12 nH, respectively. The total increase of about 0.24 nH demonstrates that the measurement arrangement changes the switching circuit. Inter-cell mutual coupling is small in the existing model, with a maximum reported coefficient of 0.034. Additional bypass capacitors give diminishing inductance reduction beyond the first few close to each cell.')
p('The four-device preliminary loss estimate is approximately 18 W per phase leg: 12.5 W conduction, 5.0 W switching, and a residual approximately 0.5 W from other terms and rounding in the existing calculation. Conduction therefore represents about 69% of the reported total. These are fundamental-period loss estimates for the assumed load and cooling model; peak-current DPT energies must not be multiplied by switching frequency and presented as the sinusoidal operating loss. The approximately 69 °C junction estimate uses the existing 7.6 K/W device-to-heatsink path and 0.65 K/W heatsink assumption and requires thermal validation.')
table(['Simulation configuration','High-side peak VDS','Other result','Interpretation'],[
['Operating board, Rg,on = 2 Ω','91 V','Eon ≈ 27.3 µJ per device','Above 90 V target'],['Operating board, Rg,on = 3 Ω','86 V','Eon ≈ 32.8 µJ per device','Nominal voltage target met'],['DPT with extracted gates, Rg,on = 2.5 Ω','105 V','di/dt ≈ 49 A/ns','Exceeds 100 V rating'],['DPT adjusted drive, Rg,on = 3.5 Ω and −1.9 V off','93 V','Eon ≈ 31.2 µJ per device','Below rating but misses 90 V target']])
doc.add_paragraph('Table V. Existing extracted-network DPT simulations at approximately the peak phase current. Gate-resistance values follow the respective existing configurations and include shared resistance where applicable.','Caption')
p('The extracted gate network is a major deviation from the original lumped assumption of 1.7 nH. Although symmetric, its larger common inductance changes switching speed and ringing. At zero off-state bias, the operating-board model reports a gate bump of about 1.2 V, above the 0.8 V minimum threshold [8]. Increasing turn-on resistance reduces drain stress but leaves the reported bump near 1.1 V, so this adjustment alone does not establish immunity to false turn-on. The adjusted DPT model uses −1.9 V off bias, giving an approximately −0.5 V held-off gate peak and −3.4 V negative ringing. Negative bias is a response to this gate network, not a universal requirement for EPC2361.')
p('The next iteration should shorten the common gate path and improve the local sink return before relying on additional negative bias. The remaining DPT overshoot may be reduced by further controlled turn-on, reduced shunt interconnect inductance, or a revised submodule voltage, with switching and reverse-conduction losses reassessed. The present evidence supports four-way paralleling conditionally; it does not yet demonstrate that every board configuration satisfies the full margin target.')

page('7 DPT platform and simulation evidence')
fig(ROOT/'icpe/img/boards.png','Fig. 5. Existing board renderings and extraction geometry. The DPT variant includes individual cell-current measurement; the operating variant removes the measurement shunts.',6.3)
fig(ROOT/'icpe/img/q3d_spice_2x2.png','Fig. 6. Existing extracted-network switching simulations. These waveforms are simulation evidence and are not experimental DPT measurements.',6.3)
p('The DPT board implements the same four-cell principle with a coaxial shunt in each cell to resolve branch currents. An inductive load establishes the test current during the first pulse; the following turn-off and second turn-on expose drain overshoot, held-off gate disturbances, and dynamic sharing. The model includes the measurement paths because the shunts materially alter the commutation inductance. Comparison to the operating board therefore uses its own extracted network rather than treating the DPT fixture as electrically transparent.')
p('For the existing operating-board model, the four cell currents at turn-off span approximately 77–81 A and the reported turn-on energy spread is about 2%. The adjusted DPT model gives approximately 78–80 A. Its rounded energy-spread output is 0%, which should be interpreted as agreement at the available reporting resolution rather than exact equality. These results support the symmetric layout under matched model assumptions. They do not cover manufacturing spread or establish experimental sharing.')

page('8 Experimental validation and conclusion')
p('The full paper is planned to include experimental DPT results on the test board. Measurements will compare simulated and observed drain overshoot, individual cell currents, held-off gate voltage, and switching energies at the same operating points and gate settings. The comparison will include the probe and shunt connections, the effective test-loop inductance, and the timing alignment used for voltage-current energy integration. Tests over increasing current and bus voltage will establish the usable operating envelope before the nominal 75 V and peak-current point is claimed as experimentally verified.')
p('Particular attention will be given to the 93 V simulated DPT peak, the gate-voltage margin, and the difference between the instrumented and operating layouts. Device-terminal gate probing and channel deskew are needed to distinguish real gate disturbances from measurement artifacts. The model will be updated from the measured waveforms, and any required gate or layout revision will be reflected in the reported loss and voltage margins. Continuous operation and thermal validation are subsequent work beyond the DPT evidence presented here.')
h('Conclusion')
p('A phase-leg building block for a 1200 V, 300 kW SPB traction drive has been developed through linked voltage-class selection, parallel-count sizing, symmetric gate-drive design, and extracted-network simulation. The existing device survey selects the 100 V EPC2361 using RDS(on) Coss,tr, while the loss and overshoot study motivates four parallel transistors per switch. Two identical driver blocks provide one output channel per pair of parallel devices, with Kelvin returns and symmetric branches to reduce dynamic imbalance. The KiCad–Q3D–LTspice iteration identifies both the low power-loop inductance and the larger-than-assumed common gate inductance. Adjusted turn-on gives an 86 V simulated operating-board peak, while the instrumented DPT board remains at 93 V and requires additional improvement to meet the 90 V target. Experimental DPT correlation will be provided in the full paper.')
h('References')
refs=[
'[1] S. Norrga, L. Jin, O. Wallmark, A. Mayer, and K. Ilves, “A novel inverter topology for compact EV and HEV drive systems,” Proc. IECON, 2013, pp. 6590–6595.',
'[2] L. Jin, S. Norrga, O. Wallmark, and M. Nikouei Harnefors, “Control and modulation of the stacked polyphase bridges inverter,” Proc. IEEE ECCE, 2014.',
'[3] L. Jin, S. Norrga, O. Wallmark, and N. Apostolopoulos, “Modulation and power losses of a stacked polyphase bridge converter,” IEEE J. Emerg. Sel. Topics Power Electron., vol. 5, no. 1, pp. 409–418, 2017.',
'[4] M. Nikouie, O. Wallmark, L. Jin, L. Harnefors, and H.-P. Nee, “DC-link stability analysis and controller design for the stacked polyphase bridges converter,” IEEE Trans. Power Electron., vol. 32, no. 2, pp. 1666–1674, 2017.',
'[5] G. Rohner, J. Huber, S. Mirić, and J. W. Kolar, “Comparative evaluation of three-phase three-level flying capacitor and stacked polyphase bridge GaN inverter systems for integrated motor drives,” Electronics, vol. 13, art. 1259, 2024. doi:10.3390/electronics13071259.',
'[6] J. Lu and D. Chen, “Paralleling GaN E-HEMTs in 10 kW–100 kW systems,” Proc. IEEE APEC, 2017, pp. 3049–3056.',
'[7] M. T. Tran et al., “A high-performance GaN power module with parallel packaging for high-current and low-voltage traction inverter applications,” IEEE J. Emerg. Sel. Topics Power Electron., vol. 13, no. 1, pp. 1188–1209, 2025.',
'[8] Efficient Power Conversion, EPC2361 datasheet, revised Feb. 20, 2026. https://epc-co.com/epc/Portals/0/epc/documents/datasheets/EPC2361_datasheet.pdf',
'[9] Infineon Technologies, 2EDF7275K product information and datasheet. https://www.infineon.com/part/2EDF7275K']
for r in refs:
 pa=p(r); pa.paragraph_format.space_after=Pt(4)
 for run in pa.runs: run.font.size=Pt(9)
doc.core_properties.title='Building Block Design for Stacked Polyphase Bridge Converters for High Voltage Traction Drives'
doc.core_properties.author='Enes Ayaz'
doc.core_properties.subject='ICPE conference manuscript draft'
doc.save(OUT)
print(OUT)
print('Paragraphs',len(doc.paragraphs),'Tables',len(doc.tables))

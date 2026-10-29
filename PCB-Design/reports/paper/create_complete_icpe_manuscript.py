from pathlib import Path
from copy import deepcopy
import re,json,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

P=Path(r'C:\Github\KTH\PCB-Design\reports\paper');A=P/'icpe';QA=A/'complete_manuscript_qa'
SRC=A/'WrittenByHand/ICPE_Template.docx'
OUT=P/'SPB_ICPE_complete_manuscript_2026-10-29.docx'
source=Document(SRC);sp=[q.text.strip() for q in source.paragraphs]
QA.mkdir(exist_ok=True)
(QA/'artifact.md').write_text(f'''# Complete ICPE manuscript contract
Reference content: {SRC}
SHA256: {hashlib.sha256(SRC.read_bytes()).hexdigest()}
Reference render: manual_reference.pdf (10 pages), reference_contact.png.
Reference roles: Times New Roman body, numbered centered section headings, italic lettered subsections, centered title, single-column anonymous paper. New manuscript replaces all body content and template placeholders; original file stays unchanged. The user authorized consolidation and remaining sections.
Page system: A4, 3 cm top and 2.5 cm other margins, single column; explicit venue instructions override inconsistent sample section geometry. Body and captions 10 pt; title 16 pt. Page field in footer. Expanded working manuscript, not a five-page submission.
Components: eight figures at 16 cm width, three IEEE-style tables without vertical rules, five editable Word equations, eleven numbered references. Black heading text. No author or affiliation.
Content: manual sections I–IV edited for coherence, corrected D-FOM and numbering; completed sections V–VII. Only simulated results are reported, DPT experiments pending. Existing manuscripts and figure masters preserved.
Rendering: packaged renderer unavailable because LibreOffice is absent; Word PDF export and Poppler page images provide visual verification.
''',encoding='utf-8')
# Reuse the retained manual template styles; clear only the content being replaced.
doc=Document(SRC);body=doc._element.body
for el in list(body):
 if el.tag!=qn('w:sectPr'):body.remove(el)
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(3);sec.bottom_margin=Cm(2.5);sec.left_margin=sec.right_margin=Cm(2.5)
sec.header_distance=Cm(1.2);sec.footer_distance=Cm(1.2)
for t in ['w:headerReference','w:footerReference','w:pgNumType']:
 for el in sec._sectPr.findall(qn(t)):sec._sectPr.remove(el)
cols=sec._sectPr.find(qn('w:cols'))
if cols is None:cols=OxmlElement('w:cols');sec._sectPr.append(cols)
cols.set(qn('w:num'),'1');cols.set(qn('w:space'),'0')
for st in doc.styles:
 if st.type==1:
  st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
  for el in st.element.xpath('.//w:pBdr | .//w:framePr | .//w:numPr'):el.getparent().remove(el)
  for el in st.element.xpath('.//w:rFonts'):
   for key in list(el.attrib):
    if 'Theme' in key:el.attrib.pop(key)
  st.paragraph_format.page_break_before=False
normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=1;normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.first_line_indent=Cm(.35);normal.paragraph_format.widow_control=True
for n in ['Title','Heading 1','Heading 2','Caption']:
 st=doc.styles[n];st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.line_spacing=1
st=doc.styles['Title'];st.font.size=Pt(16);st.font.bold=True;st.paragraph_format.space_after=Pt(9)
st=doc.styles['Heading 1'];st.font.size=Pt(10);st.font.bold=True;st.font.italic=False;st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
st=doc.styles['Heading 2'];st.font.size=Pt(10);st.font.italic=True;st.font.bold=False;st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(3);st.paragraph_format.keep_with_next=True
st=doc.styles['Caption'];st.font.size=Pt(10);st.font.bold=False;st.paragraph_format.space_after=Pt(6)
def p(t='',style=None):
 q=doc.add_paragraph(t,style);q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 q.paragraph_format.keep_together=False
 return q
def h(t):
 q=p(t,'Heading 1');q.alignment=1

def sub(t):p(t,'Heading 2')
def fig(name,caption):
 q=p();q.alignment=1;q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.keep_with_next=True;q.paragraph_format.space_after=Pt(0)
 q.add_run().add_picture(str(A/(name+'.png')),width=Cm(16))
 q=p(caption,'Caption');q.paragraph_format.keep_together=True

def table(num,title,headers,rows,widths):
 q=p('TABLE '+num+'\n'+title.upper());q.alignment=1;q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for rr in rows:
  for c,v in zip(t.add_row().cells,rr):c.text=str(v)
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=1
   for q in c.paragraphs:
    q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.space_before=Pt(2);q.paragraph_format.space_after=Pt(2);q.paragraph_format.keep_with_next=i<len(t.rows)-1;q.alignment=0 if j==0 else 1
    for r in q.runs:r.font.size=Pt(10);r.bold=i==0
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','left','right','insideH','insideV']:
  z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');z.set(qn('w:sz'),'6');borders.append(z)
 t._tbl.tblPr.append(borders)
 for c in t.rows[0].cells:
  bb=OxmlElement('w:tcBorders');z=OxmlElement('w:bottom');z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');bb.append(z);c._tc.get_or_add_tcPr().append(bb)
 return t
# Native editable OMML equations.
def mr(t,plain=False):
 z=OxmlElement('m:r');rpr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p' if plain else 'i');rpr.append(sty);z.append(rpr);v=OxmlElement('m:t');v.text=t;z.append(v);return z
def sm(a,b):
 z=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(a));ss=OxmlElement('m:sub');ss.append(mr(b,True));z.extend([e,ss]);return z
def fraction(a,b):
 z=OxmlElement('m:f')
 for tag,seq in [('m:num',a),('m:den',b)]:
  e=OxmlElement(tag)
  for v in seq:e.append(v)
  z.append(e)
 return z
def sqrt(a):
 z=OxmlElement('m:rad');pr=OxmlElement('m:radPr');v=OxmlElement('m:degHide');v.set(qn('m:val'),'1');pr.append(v);z.append(pr);z.append(OxmlElement('m:deg'));e=OxmlElement('m:e')
 for v in a:e.append(v)
 z.append(e);return z
def eq(no,seq):
 q=p();q.alignment=1;q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.space_before=Pt(3);q.paragraph_format.space_after=Pt(6);q.paragraph_format.keep_together=True
 z=OxmlElement('m:oMath')
 for v in seq:z.append(v)
 q._p.append(z);q.add_run('     ('+str(no)+')')

q=p('Topic 17 — Power Electronics for Transportation Electrification');q.paragraph_format.first_line_indent=Pt(0)
for r in q.runs:r.italic=True
q=p('Parasitic Aware Building Block Design for Stacked Polyphase Bridge Traction Inverters','Title');q.alignment=1
abstract='The stacked polyphase bridge (SPB) introduces a voltage-partitioning degree of freedom into high-voltage traction inverters: the number of series-connected submodules determines the required semiconductor voltage class. However, voltage partitioning does not reduce the phase current for a fixed drive operating point, making device paralleling and dynamic current sharing central implementation constraints. This paper presents a hardware-constrained design procedure for an SPB phase-leg building block. A device figure of merit based on on-state resistance and output charge identifies the 100 V EPC2361 as the preferred candidate among the surveyed devices. Single-cell simulations and an electrical and thermal comparison support four parallel transistors per switch position for a candidate arrangement of sixteen submodules. Current-sharing sensitivity informs a symmetric cell arrangement with one driver output channel per transistor pair. An iterative KiCad–Ansys Q3D–LTspice workflow evaluates the routed power and gate networks. The operating and instrumented double-pulse test boards have evaluated commutation-loop inductances of 0.41 nH and 0.65 nH, respectively. The extracted test-board gate network requires revised drive settings; the evaluated peak drain–source voltage remains approximately 93 V, above the 90 V design target. The study connects SPB voltage partitioning to the physical constraints governing the use of low-voltage devices at traction-current levels. Experimental double-pulse tests are planned for the full paper.'
q=p('Abstract—'+abstract);q.paragraph_format.first_line_indent=Pt(0)
for r in q.runs:r.bold=True
q=p('Index Terms—Gallium nitride, gate drive, parallel transistors, parasitic inductance, stacked polyphase bridge, traction inverter.');q.paragraph_format.first_line_indent=Pt(0)
h('I. Introduction')
p(sp[7])
p(sp[11]);q=sp[13];split=q.index('The central challenge');p(q[:split].strip());p(q[split:]);p(sp[15])
fig('SPB_introduction_topology','Fig. 1. (a) SPB traction-drive topology with N series-connected submodules, each supplying an isolated three-phase winding set of a common machine. (b) Three-phase submodule with the studied phase leg highlighted. Each switch symbol represents four parallel GaN transistors; the diode denotes the equivalent reverse-conduction path.')
h('II. Traction Requirements and Initial Device Selection')
sub('A. Traction Drive Requirements')
p('The building block is designed for the traction operating point and constraints summarized in Table I. For balanced operation with equal submodule voltages, the total active power is expressed as')
eq(1,[mr('P = '),fraction([mr('3 M '),sm('V','dc'),mr(' '),sm('I','ph,rms'),mr(' cos ',True),mr('φ')],[mr('2'),sqrt([mr('2')])])])
p('Here, Vdc is the total dc-link voltage, Iph,rms is the rms phase current, and cosφ is the power factor. The modulation index M is defined such that the fundamental phase-to-neutral voltage peak of each submodule is M Vm/2. The specified operating point requires approximately 228 A rms and 322 A peak phase current.')
table('I','Design operating point and preliminary constraints',['Quantity','Symbol','Value'],[
['Total active power','P','300 kW'],['Total dc-link voltage','Vdc','1200 V'],['Modulation index','M','1.15'],['Power factor','cosφ','0.9'],['Switching frequency','fsw','50 kHz'],['Candidate submodule count','N','16'],['Nominal submodule voltage','Vm','75 V'],['Peak drain–source voltage target','VDS,pk','90 V'],['Junction-temperature limit','Tj,max','125 °C'],['Ambient temperature','Ta','40 °C']],[8.8,2.2,5])
p('The submodule count does not appear explicitly in (1). Increasing the number of series-connected submodules reduces the voltage supplied to each winding set while increasing the number of winding sets contributing to the total power. Consequently, voltage partitioning does not reduce the phase-current requirement under the stated operating conditions. The selected devices must therefore provide both an appropriate blocking-voltage rating and sufficient current-handling capability.')
sub('B. Semiconductor Voltage Class Comparison')
p('Voltage partitioning allows the semiconductor voltage class to be selected together with the submodule count. This choice requires a comparison of conduction and switching characteristics rather than on-state resistance alone. Fig. 2 compares the typical on-state resistance RDS(on) at 25 °C and the charge-equivalent output capacitance Coss,Q of the surveyed GaN devices. The latter is calculated as')
eq(2,[sm('C','oss,Q'),mr(' = '),fraction([sm('Q','oss')],[sm('V','test')])])
p('where Qoss is the output charge at the datasheet test voltage Vtest. This quantity accounts for the charge accumulated over the voltage excursion, whereas differential output capacitance describes the capacitance at a particular voltage. Following the charge-based screening criterion in [9], the device figure of merit is defined as')
eq(3,[mr('D-FOM = ',True),fraction([mr('1')],[sqrt([sm('R','DS(on)'),mr(' '),sm('C','oss,Q')])])])
p('A higher D-FOM indicates a more favorable combination of conduction resistance and output charge. In the simplified hard-switching model underlying this criterion, optimizing semiconductor area gives minimum losses proportional to the square root of the resistance–capacitance product. Fig. 2 therefore uses constant-product contours, with lower τ = RDS(on) Coss,Q corresponding to higher D-FOM.')
fig('GaN_device_voltage_class_comparison','Fig. 2. Typical on-state resistance at 25 °C versus charge-equivalent output capacitance of the surveyed GaN devices, distinguished by voltage class. Dashed lines indicate constant resistance–capacitance products τ = RDS(on) Coss,Q. Lower τ corresponds to higher D-FOM. Charge-equivalent capacitance is calculated at the respective datasheet test voltage.')
p('Among the surveyed devices, EPC2361 has the highest D-FOM and is selected for further investigation. Its typical on-state resistance is 0.75 mΩ [10]. The quotient of the rounded output charge of 90 nC and its 50 V test voltage gives 1800 pF, consistent with the datasheet time-related capacitance of 1796 pF. This is distinct from the differential output capacitance of 1019 pF at the same voltage.')
p('The comparison is an initial screening rather than a direct prediction of converter efficiency. Output-charge values are specified at different test voltages, and the simplified criterion does not capture temperature, gate-drive, parasitic, and current-sharing effects. The candidate is subsequently evaluated at the intended submodule voltage and traction current.')
sub('C. Initial Submodule Voltage')
p('For equal voltage partitioning, the nominal submodule voltage is')
eq(4,[sm('V','m'),mr(' = '),fraction([sm('V','dc')],[mr('N')])])
p('The candidate arrangement in Table I permits 100 V devices to be considered. The available margin must accommodate submodule-voltage variation and switching overshoot. The peak drain–source voltage target leaves a 15 V allowance above the nominal submodule voltage and a further margin below the device rating. The candidate operating point is therefore assessed through switching simulations and parasitic extraction rather than accepted from the nominal rating alone.')
h('III. Single Cell Characterization and Selection of the Parallel Count')
sub('A. Single Cell Switching Characterization')
p(sp[45].replace('an external turn-on resistance','a gate-loop inductance of 1.7 nH, an external turn-on resistance'))
p(sp[47]);p(sp[49])
fig('single_cell_characterization',sp[52])
sub('B. Loss and Thermal Assessment')
p('For n parallel transistors and ideal current sharing, the peak current per device is')
eq(5,[sm('I','dev,pk'),mr(' = '),fraction([sm('I','ph,pk')],[mr('n')])])
p(sp[66]);p(sp[67]);p(sp[68].replace('Further paralleling continues to reduce the calculated loss, although the improvement becomes progressively smaller.','Further paralleling initially reduces the calculated loss, which approaches a plateau at the higher counts.'))
fig('parallel_count_selection',sp[72])
sub('C. Selection of Four Parallel Devices')
p('Table II summarizes the preliminary comparison. Two devices exceed both the thermal and voltage limits. Three satisfy the thermal constraint but still exceed the voltage target at the compact-loop assumption. Four are the smallest tested count satisfying both constraints in that case.')
table('II','Preliminary parallel count comparison',['Parallel count n','Peak device current (A)','Leg loss (W)','Junction temperature (°C)','Peak VDS at 0.5 nH (V)','Peak VDS at 1 nH (V)'],[
['2','161','39.5','141','131','180'],['3','107','23.2','84','96','123'],['4','80.5','18.0','69','83','93'],['5','64.4','15.8','62','86','85'],['6','53.7','14.8','59','86','84']],[1.6,2.7,2,3.1,3.3,3.3])
p('Loss and junction-temperature estimates in Table II use the 0.5 nH case and the preliminary gate-drive and cooling assumptions.').paragraph_format.first_line_indent=Pt(0)
p(sp[80]);p(sp[81])
h('IV. Gate Drive Configuration and Cell Arrangement')
sub('A. Current Sharing Sensitivity')
p(sp[96].replace('Fig. 5(a) summarizes','Fig. 5(a) summarizes'))
fig('paired_gate_drive_cell_layout_v2','Fig. 5. Gate-drive design considerations: (a) simulated switching-energy imbalance under the investigated parasitic mismatches; (b) four switching cells grouped into two pairs, each served by a dual-channel driver with individual gate resistors; and (c) schematic EPC2361 pad arrangement and Kelvin-source connection, illustrating separation of the gate return from the power-source inductance. Driver supply and return connections are omitted in (b) for clarity. High-side (HS) and low-side (LS) outputs are denoted HO and LO.')
sub('B. Paired Gate Drive Configuration')
p(sp[100]);p(sp[101]);p(sp[102].replace('8 A per channel.','8 A per channel [11].'))
sub('C. Kelvin Source Connection and Cell Layout')
for i in [105,106,107]:p(sp[i])
sub('D. Verification of the Physical Arrangement')
p(sp[109]);p('The paired configuration is translated into a routed board and evaluated through the extraction-based workflow in Section V. The assessment determines whether the current capability and switching margin assumed during initial sizing can be retained in the physical building block.')
h('V. Closing the Design Loop with the Routed Building Block')
sub('A. Layout and Parasitic Extraction')
p('The initial sizing establishes a candidate voltage, parallel count, and drive configuration using assumed loop inductances. The routed building block must then be assessed to determine whether these assumptions are physically achievable. Fig. 6 shows the iterative KiCad–Ansys Q3D–LTspice workflow connecting conductor geometry to switching behavior.')
p('The power devices, local decoupling capacitors, and paired gate drivers are placed and routed in KiCad. The conductor geometry is evaluated in Ansys Q3D to obtain resistive and inductive networks, including mutual coupling between branches. These networks are incorporated into LTspice to assess voltage overshoot, gate disturbance, transient current sharing, and switching-energy distribution. Placement, routing, or drive settings are revised where necessary.')
fig('routed_design_workflow','Fig. 6. Iterative design and verification workflow. Extracted power and gate networks connect the routed geometry to switching simulations. The results inform revisions to layout, drive settings, and, where necessary, voltage partition or parallel count. Experimental double-pulse test (DPT) correlation is planned.')
p('Two board variants are considered: an operating board representing the intended converter building block and an instrumented DPT board. The latter includes current-measurement provisions that introduce additional commutation-path parasitics. Figs. 7(a) and 7(b) distinguish these implementations, preventing measurement-related inductance from being treated as an unavoidable feature of the operating converter.')
sub('B. Extracted Power and Gate Networks')
p('Fig. 7(c) separates the commutation-loop contributions. The operating-board loop is approximately 0.41 nH per cell, supporting the feasibility of the compact-loop assumption used in preliminary sizing. The additional conductor length and shunt contribution increase the DPT loop to approximately 0.65 nH. The capacitor contribution is included through its equivalent series inductance (ESL).')
p('The DPT gate network departs more substantially from the preliminary assumptions. Fig. 7(d) compares the individual loops with the pair-driven loops, which include the effect of the shared gate path. Symmetric branches can therefore retain substantial shared inductance. Matching alone does not establish a sufficiently low gate-loop impedance.')
fig('board_parasitic_comparison','Fig. 7. Building-block variants and evaluated parasitics: (a) operating-board rendering; (b) instrumented DPT-board rendering; (c) commutation-loop inductance including conductor, capacitor ESL, and shunt contributions; and (d) individual and pair-driven DPT gate-loop inductances. Gate-network extraction in this study applies to the DPT board.')
sub('C. Switching Assessment and Design Feedback')
p('Switching simulations are performed near the required phase-current peak, with approximately 316 A achieved in the evaluated DPT cases. Table III compares the saved configurations. At equal DPT turn-on resistance, replacing the preliminary lumped gate model with the extracted gate network increases the peak drain–source voltage. Reducing the resistance in the extracted-network case increases voltage stress further, demonstrating that the gate network materially changes the available switching margin.')
table('III','Switching simulation comparison near 316 A',['Board','Gate model','Rg,on (Ω)','Off bias (V)','Peak VDS (V)','Mean Eon per device (µJ)'],[
['Operating','Preliminary','2.0','0','91','27.3'],['Operating','Preliminary','3.0','0','86','32.8'],['DPT','Preliminary','3.5','−1.9','90','34.1'],['DPT','Extracted','2.5','−1.9','105','24.3'],['DPT','Extracted','3.5','−1.9','93','31.2']],[2.4,3,2.2,2.3,2.6,3.5])
p('The operating-board cases include the evaluated power path but retain the preliminary gate model. The DPT extracted cases include both power and gate networks. Switching-energy values are device averages for the evaluated turn-on event, not fundamental-period average losses.').paragraph_format.first_line_indent=Pt(0)
p('Fig. 8(a) illustrates the associated drain-voltage transients. Slower turn-on reduces stress at the cost of greater switching energy, so drive resistance must be selected through a combined voltage and loss assessment. The evaluated DPT configuration remains below the nominal device rating but above the project voltage target.')
p('Fig. 8(b) compares the held-off high-side gate voltage for two board and drive configurations. The zero-bias operating-board case rises above the minimum specified gate threshold, indicating insufficient margin against unintended turn-on. Negative turn-off bias suppresses the positive excursion in the extracted-network DPT case. Because both the board network and bias differ, the comparison illustrates complete configurations rather than isolating the effect of bias alone.')
fig('extracted_network_switching','Fig. 8. Simulated switching at 75 V and approximately 316 A: (a) worst-cell high-side drain–source voltage during low-side turn-on for the indicated board, gate-network, and resistance configurations; and (b) cell-1 held-off high-side gate voltage. Reference lines denote the 90 V drain-voltage target and the 0.8 V minimum gate threshold. Operating-board cases retain the preliminary gate model; DPT cases include the extracted gate network.')
p('The evaluated negative-bias DPT case has a maximum held-off gate excursion of approximately −0.5 V and negative ringing reaching approximately −3.4 V over the assessed transitions. Both positive and negative gate excursions must be checked against the device limits; the minimum gate rating is −4 V [10].')
sub('D. Implications for the Building Block')
p('The power-network results support a compact commutation path, while the gate-network results identify shared routing inductance as a remaining constraint. The evaluated extracted-network DPT case is therefore a candidate requiring refinement rather than a completed validation against all targets. Shorter shared gate paths and driver output stages closer to each pair are the next layout objectives.')
p('Updated switching energies must be used to revise the preliminary loss and thermal estimates. The operating-board simulations must also be repeated with their corresponding extracted gate network before its switching margin can be established. This feedback connects the physical building block to the initial SPB voltage choice: the expected benefit of low-voltage devices depends on the power and gate networks together.')
h('VI. Experimental Validation and Scope of the Full Paper')
sub('A. Planned Double Pulse Measurements')
p('The instrumented board provides the basis for experimental DPT correlation. The planned assessment starts at reduced voltage and current before progressing toward the candidate submodule operating point. Measurements will include drain–source voltage, gate–source voltage referenced to the Kelvin terminal, total commutating current, and individual branch currents where the sensing arrangement permits.')
p('The measured voltage and current waveforms will be used to assess overshoot, ringing, transient current sharing, and switching energy. Probe delay and channel skew must be accounted for before voltage–current integration. The effects of the shunts, connections, and probe loading will be considered when comparing measurements with the instrumented-board model. Turn-on resistance and turn-off bias will be varied to examine the voltage-stress and loss tradeoff identified by simulation.')
p('Experimental results are not available in this manuscript. The full paper is intended to report correlation of the DPT model and measurements, including whether the voltage and gate-drive margins predicted by the revised layout are retained in hardware. Operating-board loss and temperature validation will require a subsequent continuous-current assessment under a defined cooling condition.')
sub('B. System Integration Boundary')
p('The present study addresses one phase-leg building block. It does not establish a globally optimum SPB submodule count or a complete traction-drive efficiency advantage. Submodule-voltage balancing, winding design, common-mode behavior, auxiliary-power overhead, and insulation across the series stack remain system-level constraints. The voltage-class comparison and preliminary loss model do not replace those evaluations.')
p('Inter-driver timing and the common interconnections between cell pairs must also be evaluated. The gate driver provides functional isolation [11]; it does not by itself establish the insulation architecture required for a submodule positioned within the full dc-link stack. These integration requirements bound the conclusions that can be drawn from the local DPT model.')
h('VII. Conclusion')
p('A hardware-constrained procedure has been presented for a phase-leg building block of an SPB traction inverter. Charge-based device screening selects the 100 V EPC2361 from the surveyed candidates. Single-cell characterization and an electrical and thermal comparison support four parallel devices per switch position as the smallest tested count meeting the preliminary constraints for the compact-loop assumption. Current-sharing sensitivity motivates paired gate drive, matched gate branches, and dedicated Kelvin returns.')
p('The routed-network assessment shows why device selection and voltage partitioning cannot be finalized independently of implementation. A compact power loop is achievable in the evaluated operating-board geometry, but the extracted DPT gate network changes the required switching settings. The remaining gap to the project voltage target requires further gate-network refinement and updated loss estimates. Experimental DPT correlation is planned for the full paper. The contribution is the traceable connection between SPB voltage choice and the physical switching constraints that determine the usefulness of low-voltage devices at traction-current levels.')
h('References')
old=Document(A/'SPB_BB_ICPE2027_realistic_design_flow_2026-10-23.docx')
refs=[q.text for q in old.paragraphs if re.match(r'^\[\d+\]',q.text)]
refs=refs[:8]+['[9] J. Azurza Anderson, G. Zulauf, J. W. Kolar, and G. Deboy, “New figure-of-merit combining semiconductor and multi-level converter properties,” IEEE Open J. Power Electron., vol. 1, pp. 322–338, 2020, doi: 10.1109/OJPEL.2020.3018220.', '[10] Efficient Power Conversion, “EPC2361 eGaN FET datasheet,” rev. July 20, 2026. [Online]. Available: https://epc-co.com/epc/Portals/0/epc/documents/datasheets/EPC2361_datasheet.pdf', '[11] Infineon Technologies, “2EDF7275K data sheet,” rev. 2.8. [Online]. Available: https://www.infineon.com/part/2EDF7275K']
for r in refs:
 q=p(r);q.paragraph_format.first_line_indent=Cm(-.65);q.paragraph_format.left_indent=Cm(.65);q.paragraph_format.space_after=Pt(3);q.alignment=0
# Symbols remain editable, with true subscripts throughout prose and tables.
symbols={'Iph,rms':('I','ph,rms'),'Iph,pk':('I','ph,pk'),'Idev,pk':('I','dev,pk'),'VDS,pk':('V','DS,pk'),'RDS(on)':('R','DS(on)'),'Coss,Q':('C','oss,Q'),'Vtest':('V','test'),'Qoss':('Q','oss'),'Rg,on':('R','g,on'),'Tj,max':('T','j,max'),'Vdc':('V','dc'),'Vm':('V','m'),'fsw':('f','sw'),'VDS':('V','DS'),'Eon':('E','on'),'Ta':('T','a')}
pat=re.compile(r'(?<![A-Za-z])('+'|'.join(re.escape(k) for k in sorted(symbols,key=len,reverse=True))+r')(?![A-Za-z])')
paras=list(doc.paragraphs)+[q for t in doc.tables for row in t.rows for c in row.cells for q in c.paragraphs]
for q in paras:
 if re.match(r'^\[\d+\]',q.text):continue
 for rr in list(q.runs):
  old=rr.text
  if not pat.search(old):continue
  base=deepcopy(rr._r.rPr) if rr._r.rPr is not None else None;pos=0
  chunks=[]
  for m in pat.finditer(old):
   chunks.extend([(old[pos:m.start()],False,False),(symbols[m.group()][0],False,True),(symbols[m.group()][1],True,False)]);pos=m.end()
  chunks.append((old[pos:],False,False))
  for val,isub,ital in chunks:
   if not val:continue
   r=OxmlElement('w:r')
   if base is not None:r.append(deepcopy(base))
   rp=r.find(qn('w:rPr'))
   if rp is None:rp=OxmlElement('w:rPr');r.append(rp)
   if isub:
    z=OxmlElement('w:vertAlign');z.set(qn('w:val'),'subscript');rp.append(z)
   if ital:rp.append(OxmlElement('w:i'))
   z=OxmlElement('w:t');z.set(qn('xml:space'),'preserve');z.text=val;r.append(z);rr._r.addprevious(r)
  rr._r.getparent().remove(rr._r)
footer=sec.footer.paragraphs[0];footer.alignment=1;footer.paragraph_format.first_line_indent=Pt(0)
r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='Parasitic Aware Building Block Design for Stacked Polyphase Bridge Traction Inverters';doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.subject='Expanded ICPE working manuscript with simulated results'
settings=doc.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
doc.save(OUT)
print(OUT)
print('Words:',sum(len(q.text.split()) for q in doc.paragraphs),'Figures:',len(doc.inline_shapes),'Tables:',len(doc.tables),'Equations:',len(doc.element.xpath('.//m:oMath')))

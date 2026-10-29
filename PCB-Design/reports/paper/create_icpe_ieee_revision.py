"""Create a separate ICPE digest from the venue template and saved design data."""
from pathlib import Path
import json, re, sys, hashlib, subprocess, math
import numpy as np
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

P=Path(__file__).resolve().parent; R=P.parents[1]
A=P/'icpe'/'ieee_revision_assets'; A.mkdir(exist_ok=True)
OUT=P/'icpe'/'SPB_BB_ICPE2027_IEEE_revised_2026-10-23.docx'
TPL=P/'ICPE_Template.docx'
deps=Path(sys.executable).parent.parent
pop=deps/'native/poppler/Library/bin/pdftoppm.exe'
if not pop.exists(): pop=Path('C:/Users/enesa/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe')
pdfmetrics.registerFont(TTFont('TimesLocal','C:/Windows/Fonts/times.ttf'))
pdfmetrics.registerFont(TTFont('TimesItalicLocal','C:/Windows/Fonts/timesi.ttf'))
INK='#111111'; TEAL='#018789'; WINE='#A41C59'; OCHRE='#C38714'; GREY='#898781'
COL=[INK,TEAL,WINE,OCHRE]; STY=[[],[4,2],[5,2,1,2],[1,2]]
W=6.3*72
def text(c,x,y,t,size=8,align='left',col=INK):
 c.setFillColor(col); c.setFont('TimesLocal',size)
 getattr(c,{'left':'drawString','center':'drawCentredString','right':'drawRightString'}[align])(x,y,t)
def line(c,pts,col=INK,dash=None,lw=1):
 c.setStrokeColor(col); c.setLineWidth(lw); c.setDash(dash or [])
 z=c.beginPath(); z.moveTo(*pts[0])
 for pt in pts[1:]: z.lineTo(*pt)
 c.drawPath(z); c.setDash([])
def finish(c,name):
 c.save(); subprocess.run([str(pop),'-singlefile','-r','600','-png',str(A/(name+'.pdf')),str(A/name)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
def axes(c,box,xlim,ylim,xt,yt,xlabel,ylabel):
 x,y,w,h=box
 def tr(a,b): return x+(a-xlim[0])/(xlim[1]-xlim[0])*w,y+(b-ylim[0])/(ylim[1]-ylim[0])*h
 for b in yt:
  xx,yy=tr(xlim[0],b); line(c,[(x,yy),(x+w,yy)],'#E2E0DA',lw=.3); text(c,x-5,yy-2.5,str(b),align='right')
 for a in xt:
  xx,yy=tr(a,ylim[0]);line(c,[(xx,y),(xx,y-3)],lw=.5);text(c,xx,y-13,str(a),align='center')
 line(c,[(x,y+h),(x,y),(x+w,y)],lw=.5)
 text(c,x+w/2,y-26,xlabel,align='center')
 c.saveState();c.translate(x-29,y+h/2);c.rotate(90);text(c,0,0,ylabel,align='center');c.restoreState()
 return tr

# One common figure style, with vector masters and 600 dpi Word images.
rf=json.loads((R/'reports/slides/data/revb_flow_numbers.json').read_text())
c=canvas.Canvas(str(A/'selection.pdf'),pagesize=(W,166))
t=axes(c,(35,39,100,105),(.5,4.5),(0,8),[1,2,3,4],[0,2,4,6,8],'Voltage class (V)','D-FoM (ps)')
for i,v in enumerate([1.35,4.14,6.88,4.90],1):
 x0,y0=t(i-.26,0);x1,y1=t(i+.26,v);c.setFillColor(TEAL);c.rect(x0,y0,x1-x0,y1-y0,stroke=0,fill=1)
 text(c,*t(i,v+.22),f'{v:.2f}',align='center')
for i,s in enumerate(['100','200','350','650'],1):
 xx,_=t(i,0);c.setFillColor('white');c.rect(xx-7,22,14,11,stroke=0,fill=1);text(c,xx,26,s,align='center')
text(c,85,3,'(a)',align='center')
t=axes(c,(190,39,100,105),(2,8),(0,45),[2,4,6,8],[0,15,30,45],'Devices per switch','Leg loss (W)')
nr=[v for v in rf['npar']['0.5nH'] if v['n']>=2]
line(c,[t(v['n'],v['tot']) for v in nr],TEAL)
line(c,[t(v['n'],v['cond']) for v in nr],INK,[4,2])
text(c,220,132,'Total',col=TEAL);text(c,220,120,'Conduction',col=INK)
text(c,240,3,'(b)',align='center')
t=axes(c,(345,39,100,105),(2,8),(75,185),[2,4,6,8],[80,100,140,180],'Devices per switch','Peak drain voltage (V)')
for key,col,dash in [('0.5nH',INK,[]),('1nH',WINE,[4,2])]:
 line(c,[t(v['n'],v['vpk']) for v in rf['npar'][key] if v['n']>=2],col,dash)
line(c,[t(2,90),t(8,90)],GREY,[1,2],.6)
text(c,384,132,'1 nH',col=WINE);text(c,384,120,'0.5 nH')
text(c,395,3,'(c)',align='center')
finish(c,'selection')

# Gate-drive architecture and extraction loop.
c=canvas.Canvas(str(A/'workflow.pdf'),pagesize=(W,155))
def box(x,y,w,h,ls):
 c.setStrokeColor(INK);c.setLineWidth(.6);c.setFillColor('white');c.rect(x,y,w,h,fill=1,stroke=1)
 for j,l in enumerate(ls): text(c,x+w/2,y+h/2+((len(ls)-1)/2-j)*11-2,l,align='center')
def arrow(a,b):
 line(c,[a,b],lw=.65);dx=b[0]-a[0];dy=b[1]-a[1];ang=math.atan2(dy,dx);L=4
 pts=[b,(b[0]-L*math.cos(ang-.5),b[1]-L*math.sin(ang-.5)),(b[0]-L*math.cos(ang+.5),b[1]-L*math.sin(ang+.5))]
 c.setFillColor(INK);z=c.beginPath();z.moveTo(*pts[0]);z.lineTo(*pts[1]);z.lineTo(*pts[2]);z.close();c.drawPath(z,fill=1,stroke=0)
box(3,105,125,40,['SPB requirements','Device class and parallel count'])
box(160,105,125,40,['Symmetric KiCad layout','Power cells and gate returns'])
box(317,105,133,40,['Ansys Q3D','Power and gate R and L'])
arrow((128,125),(160,125));arrow((285,125),(317,125))
box(317,33,133,43,['LTspice double-pulse test','Sharing, energy, drain stress','and held-off gate voltage'])
arrow((383,105),(383,76))
box(160,33,125,43,['Check design targets','Revise layout and gate drive'])
arrow((317,55),(285,55));arrow((222,76),(222,105))
box(3,33,125,43,['DPT board','Experimental correlation','in the full paper'])
arrow((160,55),(128,55))
text(c,W/2,8,'One dual-channel driver per two cells; one output channel per two parallel transistors.',align='center')
finish(c,'workflow')

# Use the actual latest raw files, including the extracted DPT gate network.
def read_raw(path):
 data=path.read_bytes(); mark='Binary:\n'.encode('utf-16-le'); off=data.index(mark)+len(mark); hdr=data[:off].decode('utf-16-le')
 nv=int(re.search(r'No. Variables:\s*(\d+)',hdr).group(1));npnt=int(re.search(r'No. Points:\s*(\d+)',hdr).group(1))
 names=re.findall(r'^\t\d+\t(\S+)\t',hdr,re.M)
 dt=np.dtype([('t','<f8')]+[(str(i),'<f4') for i in range(1,nv)])
 a=np.frombuffer(data,dt,count=npnt,offset=off); vals={names[i].lower():np.abs(a['t']) if i==0 else a[str(i)] for i in range(nv)}
 return vals
raw={k:read_raw(R/'simulation/bb_spice'/('REVB_Q3D_'+k+'.raw')) for k in ['load','load_rg3','dpt','dpt_rg2.5']}
T2=.5e-6+10.667e-6+1e-6
def wave(c,tr,z,signal,col,dash):
 tt=z['time']; ix=np.flatnonzero((tt>T2-5e-9)&(tt<T2+45e-9)); ix=ix[::max(1,len(ix)//850)]
 yy=signal(z);line(c,[tr((tt[i]-T2)*1e9,float(yy[i])) for i in ix],col,dash)
c=canvas.Canvas(str(A/'switching.pdf'),pagesize=(W,197))
t=axes(c,(35,46,110,126),(-5,45),(-10,150),[0,20,40],[0,50,100,150],'Time (ns)','Cell drain current (A)')
for j in range(4):wave(c,t,raw['dpt'],lambda z,j=j:z[f'i(vil{j+1})'],COL[j],STY[j])
text(c,90,183,'DPT board',align='center');text(c,90,8,'(a)',align='center')
t=axes(c,(190,46,110,126),(-5,45),(0,115),[0,20,40],[0,50,100],'Time (ns)','High-side drain voltage (V)')
for key,col,dash in [('load',INK,[]),('load_rg3',TEAL,[4,2]),('dpt',WINE,[5,2,1,2]),('dpt_rg2.5',OCHRE,[1,2])]:
 z=raw[key]; peaks=[np.max((z[f'v(dh{j})']-z[f'v(qh_s_{j})'])[(z['time']>T2)&(z['time']<T2+150e-9)]) for j in range(1,5)];j=int(np.argmax(peaks))+1
 wave(c,t,z,lambda z,j=j:z[f'v(dh{j})']-z[f'v(qh_s_{j})'],col,dash)
line(c,[t(-5,90),t(45,90)],GREY,[1,2],.6)
text(c,245,183,'Worst cell in each case',align='center');text(c,245,8,'(b) Configurations in Table III',align='center')
t=axes(c,(345,46,100,126),(-5,45),(-4,2),[0,20,40],[-4,-2,0,2],'Time (ns)','Held-off gate voltage (V)')
for key,col,dash in [('load',INK,[]),('dpt',WINE,[5,2,1,2])]:
 z=raw[key]; j=1;wave(c,t,z,lambda z,j=j:z[f'v(gh{j})']-z[f'v(qh_s_{j})'],col,dash)
line(c,[t(-5,.8),t(45,.8)],GREY,[1,2],.6);text(c,395,183,'Cell 1',align='center');text(c,395,8,'(c) Gate-to-source voltage',align='center')
finish(c,'switching')

# Composite board drawing with legible consistent panel labels, image panels at 600 dpi.
c=canvas.Canvas(str(A/'boards.pdf'),pagesize=(W,134))
for x,y,w,h,path,label in [(0,18,145,113,P/'figures/revb/r3d_dpt_top.png','(a)'),(153,18,145,113,P/'figures/revb/r3d_load.png','(b)'),(306,18,147,113,P/'figures/revb/q3d_loop_dpt.png','(c)')]:
 c.drawImage(str(path),x,y,width=w,height=h,preserveAspectRatio=True,anchor='c',mask='auto');text(c,x+w/2,4,label,align='center')
finish(c,'boards')

# The template is retained; the new copy replaces its example material.
contract=f'''# ICPE reference contract\nReference: {TPL}\nSHA256: {hashlib.sha256(TPL.read_bytes()).hexdigest()}\nA4 single-column anonymous digest, maximum five pages, minimum body size 10 pt.\nTemplate prose specifies top margin 3 cm and other margins 2.5 cm, overriding inconsistent sample section geometry.\nReuse Times body, centered 16 pt title, numbered sections, abstract, captions, and references.\nReplace all demonstration text and example objects. Remove guideline pages. Add page numbers as explicitly required.\nThe local IEEE skill supplies no-vertical-rule tables, first-use definitions, numerical interpretation, cross-references, and uniform figure styling.\nVector masters are exported at the venue text width, with 600 dpi images embedded in Word.\n'''
(A/'artifact.md').write_text(contract,encoding='utf8')
doc=Document(TPL)
body=doc._element.body
for el in list(body):
 if el.tag!=qn('w:sectPr'):body.remove(el)
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(3);sec.bottom_margin=Cm(2.5);sec.left_margin=sec.right_margin=Cm(2.5)
cols=sec._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'1');cols.set(qn('w:space'),'0')
for x in sec._sectPr.findall(qn('w:footerReference')):sec._sectPr.remove(x)
for st in doc.styles:
 if st.type==1:
  st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
  for x in st.element.xpath('.//w:pBdr | .//w:framePr | .//w:numPr'):x.getparent().remove(x)
  for x in st.element.xpath('.//w:rFonts'):
   for k in list(x.attrib):
    if 'Theme' in k:x.attrib.pop(k)
normal=doc.styles['Normal'];normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=1;normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.first_line_indent=Cm(.35)
for n in ['Title','Heading 1','Heading 2','Caption']:
 doc.styles[n].paragraph_format.first_line_indent=Pt(0)
doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 1'].font.size=Pt(10);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(7);doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
doc.styles['Heading 2'].font.size=Pt(10);doc.styles['Heading 2'].font.italic=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(5);doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
doc.styles['Caption'].font.size=Pt(10);doc.styles['Caption'].font.bold=False
doc.styles['Caption'].paragraph_format.space_before=Pt(2);doc.styles['Caption'].paragraph_format.space_after=Pt(5)
def p(t,style=None):
 pa=doc.add_paragraph(t,style);pa.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 for x in pa._p.xpath('./w:pPr/w:framePr'):x.getparent().remove(x)
 return pa
def h(t):
 pa=doc.add_paragraph(t,'Heading 1');pa.alignment=1
def sub(t):doc.add_paragraph(t,'Heading 2')
def pg():doc.add_page_break()
def fig(name,cap):
 pa=p('');pa.paragraph_format.first_line_indent=Pt(0);pa.alignment=1;pa.paragraph_format.space_after=Pt(0)
 pa.add_run().add_picture(str(A/(name+'.png')),width=Cm(16));pa.paragraph_format.keep_with_next=True
 cp=p(cap,'Caption');cp.paragraph_format.keep_together=True
def tab(num,title,head,rows,widths):
 pa=p('TABLE '+num+'\n'+title.upper());pa.alignment=1;pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(head));t.autofit=False
 for c,s in zip(t.rows[0].cells,head):c.text=s
 for row in rows:
  for c,s in zip(t.add_row().cells,row):c.text=str(s)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);c.vertical_alignment=1
   for pa in c.paragraphs:
    pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(1);pa.paragraph_format.space_before=Pt(1)
    for r in pa.runs:r.font.size=Pt(10)
 for col,w in zip(t.columns,widths): col.width=Cm(w)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','insideH','insideV','left','right']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');borders.append(e)
 t._tbl.tblPr.append(borders)
 for c in t.rows[0].cells:
  pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');b.append(e);pr.append(b)
def eq(txt,no):
 pa=p('');pa.paragraph_format.first_line_indent=Pt(0);pa.alignment=1
 def mr(t,plain=False):
  r=OxmlElement('m:r')
  if plain:
   rp=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');rp.append(sty);r.append(rp)
  tt=OxmlElement('m:t');tt.text=t;r.append(tt);return r
 def subm(a,b):
  z=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(a));ss=OxmlElement('m:sub');ss.append(mr(b,True));z.append(e);z.append(ss);return z
 def sq(seq):
  z=OxmlElement('m:rad');pr=OxmlElement('m:radPr');dh=OxmlElement('m:degHide');dh.set(qn('m:val'),'1');pr.append(dh);z.append(pr);z.append(OxmlElement('m:deg'));ee=OxmlElement('m:e')
  for a in seq:ee.append(a)
  z.append(ee);return z
 om=OxmlElement('m:oMath')
 if no==1:
  om.append(mr('P = '));f=OxmlElement('m:f');num=OxmlElement('m:num');den=OxmlElement('m:den')
  for a in [mr('3 M '),subm('V','dc'),mr(' '),subm('I','ph,rms'),mr(' cos ',True),mr('φ')]:num.append(a)
  den.append(mr('2'));den.append(sq([mr('2')]))
  f.append(num);f.append(den);om.append(f)
 else:
  for a in [subm('P','min'),mr(' ∝ '),subm('V','dc'),mr(' '),subm('I','ph,rms'),mr(' '),sq([subm('f','sw'),mr(' D-FoM',True)]),mr(',    D-FoM = ',True),subm('R','DS(on)'),mr(' '),subm('C','oss,tr')]:om.append(a)
 pa._p.append(om);pa.add_run('     ('+str(no)+')')

topic=p('Topic 17 — Power Electronics for Transportation Electrification');topic.paragraph_format.first_line_indent=Pt(0)
for rr in topic.runs:rr.italic=True
title=doc.add_paragraph('Building Block Design for Stacked Polyphase Bridge Converters for High Voltage Traction Drives','Title');title.alignment=1
pa=p('Abstract—A stacked polyphase bridge (SPB) enables low voltage semiconductors in a high voltage traction drive, but each submodule must carry the full phase current. This paper presents a phase-leg building block designed by jointly selecting the device voltage class, parallel count, gate drive, and layout. A device figure of merit based on on-resistance and output charge favors the 100 V EPC2361 among the surveyed candidates. Loss and voltage-stress calculations select four parallel transistors per switch. Dynamic current-sharing sensitivity motivates symmetric commutation cells and one gate-driver channel per pair of transistors. Routed power and gate networks are extracted in Ansys Q3D and inserted into LTspice double-pulse test (DPT) simulations. The operating and DPT boards have effective cell-loop inductances of 0.41 and 0.65 nH, respectively. Simulation shows that the larger common gate inductance changes the required drive settings even when branch symmetry is maintained. Slower turn-on reduces drain stress, while negative off-state bias suppresses the held-off gate rise on the instrumented board. Experimental DPT results will be included in the full paper.');pa.paragraph_format.first_line_indent=Pt(0)
for rr in pa.runs:rr.bold=True
p('Index Terms—Gallium nitride, gate drive, parallel transistors, parasitic inductance, stacked polyphase bridge, traction inverter.').paragraph_format.first_line_indent=Pt(0)
h('I. Introduction')
p('High voltage traction drives combine a large dc-link voltage with demanding current and loss requirements. A two-level inverter places the full dc-link voltage across each semiconductor switch. The stacked polyphase bridge (SPB) instead connects three-phase bridges in series on the dc side, each feeding an isolated winding set of a multiphase machine [1], [2]. The submodule voltage can therefore be selected independently of the total traction voltage. This enables low voltage devices and reduces the voltage step applied to each winding set. The resulting benefit depends on the converter and machine design, including modulation, voltage balancing, and integration [3]–[5].')
p('SPB studies have addressed modulation, dc-link stability, and system-level loss and volume [2]–[5]. Implementing a high-current submodule introduces a coupled hardware problem: reducing the device voltage class increases the submodule count, while the required phase current remains unchanged. Several low voltage transistors must then operate in parallel. Prior GaN studies identify commutation inductance, common-source coupling, and unequal gate impedances as limits to usable current capability [6]–[8]. A device-only loss comparison cannot resolve these constraints.')
p('This paper links voltage-class and parallel-count selection to symmetric gate-drive design and parasitic extraction of a phase-leg building block. Switching simulation then quantifies how power- and gate-loop parasitics change the usable voltage margin. An instrumented double-pulse test (DPT) board and an operating board distinguish the parasitics required for measurement from those of the intended converter.')
h('II. Voltage Class and Parallel Count')
sub('A. Submodule requirements')
p('Table I defines the design point. For N equally loaded submodules, the submodule voltage Vm = Vdc/N. With sinusoidal phase current, the total active power is')
eq('P = 3 M Vdc Iph,rms cos φ / (2√2).',1)
p('Here, Vdc is the total dc-link voltage, M is the modulation index, Iph,rms is the phase rms current, and φ is the displacement angle. Equation (1) makes the central design constraint explicit: increasing N reduces the device blocking voltage but does not reduce the required phase current.')
tab('I','Design point and preliminary limits',['Parameter','Value','Parameter','Value'],[
['Vdc / P','1200 V / 300 kW','Switching frequency','50 kHz'],['M / cos φ','1.15 / 0.90','Phase rms / peak current','228 A / 322 A'],['Candidate N / Vm','16 / 75 V','Drain-voltage target','90 V']], [3.3,3.1,5.5,4.1])

pg()
sub('B. Device figure of merit')
p('For n devices per switch, conduction loss scales with N Iph,rms² RDS(on)/n, whereas output-charge loss scales with N n fsw Coss,tr Vm². The switching frequency is fsw, RDS(on) is the on-resistance, and Coss,tr = Qoss/Vtest is the charge-equivalent output capacitance obtained from datasheet output charge Qoss at test voltage Vtest. Defining the device figure of merit (D-FoM) as their product, minimizing this simplified sum over n gives')
eq('Pmin ∝ Vdc Iph,rms √(fsw D-FoM),    D-FoM = RDS(on) Coss,tr.',2)
p('Equation (2) ranks candidate devices without assigning an intrinsic advantage to a voltage class. Fig. 1(a) compares the best surveyed candidate in each class: EPC2361, EPC2215, EPC2050, and IGT65R140D2, respectively. EPC2361 is selected for detailed sizing because it gives the lowest D-FoM and provides a Kelvin source connection [9]. The comparison uses typical on-resistance at room temperature and charge-equivalent capacitance at the respective datasheet test voltages. It is a screening calculation; nonlinear capacitance, junction temperature, and integer parallel counts are evaluated subsequently.')
fig('selection','Fig. 1. Device and parallel-count selection. (a) Minimum surveyed D-FoM in each voltage class (V). (b) Total semiconductor and conduction loss per leg at 0.5 nH per cell. (c) Peak drain voltage at 0.5 nH (solid) and 1 nH (dashed); the dotted line denotes the design target. The parallel-count studies use the design point in Table I.')
sub('C. Loss and voltage margin')
p('Fig. 1(b) evaluates the tradeoff between conduction loss and the charge switched by additional transistors. Switching energies are obtained from device-level DPT simulations and averaged over the fundamental current waveform. The thermal calculation iterates on-resistance and junction temperature with the assumed cooling path. A 40 °C ambient, 0.65 K/W heatsink, and 7.6 K/W junction-to-heatsink path per device give a junction estimate of 69 °C for four parallel devices. The estimate is conditional on that thermal boundary.')
p('The peak drain-to-source voltage VDS,pk is screened using VDS,pk ≈ Vm + Lloop |di/dt|, where Lloop and di/dt refer to the inductance and current slew rate of the same cell commutation path. Fig. 1(c) shows that three devices do not provide the required margin, while four satisfy the preliminary target at the lower loop inductance. More devices reduce conduction loss, but the incremental saving becomes small relative to the added gate-drive and layout complexity. Four devices per switch are therefore selected as a practical compromise. The higher-inductance case shows that this choice remains conditional on the routed loop and drive settings.')
h('III. Symmetric Cells and Gate Drive')
p('Dynamic sharing is evaluated using two parallel EPC2361 devices at 75 V and 120 A. Table II summarizes individual perturbations. Gate-loop and common-source mismatch produce the largest switching-energy imbalance. The layout must consequently match the complete gate and return impedances, rather than only the lengths of the power conductors. Kelvin returns prevent the main source-current path from becoming the gate reference.')
tab('II','Sensitivity of switching-energy sharing',['Perturbation','Energy imbalance','Layout implication'],[
['Power-path inductance +0.3 nH','Negligible','Repeated commutation cells'],['Gate loops 2 and 8 nH','19%','Equal gate branches'],['No Kelvin return; source mismatch +0.2 nH','29%','Individual Kelvin returns']], [7.0,3.1,5.9])

pg()
p('The leg contains four commutation cells, each comprising a high-side and a low-side transistor with local bypass capacitors. Two identical cells form a driver block. One dual-channel 2EDF7275K drives each block: its high-side output serves two parallel high-side transistors, and its low-side output serves the corresponding low-side pair. Two driver integrated circuits therefore serve the complete eight-transistor leg. Each output channel has 4 A source and 8 A sink capability [10]. Dividing the gates into pairs limits the charge seen by one channel and permits a symmetric star connection with individual gate resistors and Kelvin returns.')
p('Strong sink current reduces the impedance through which Miller current raises the held-off gate. It does not compensate for a long shared gate path. Point symmetry about the driver matches the two branches, while the shared path must be kept short. The implemented arrangement is therefore evaluated using both power- and gate-loop extraction. Device-parameter and driver-timing spread are not represented by geometric symmetry and remain outside the nominal sharing result.')
h('IV. Layout and Parasitic Extraction')
p('Fig. 2 gives the design procedure. Device and parallel-count selection establish initial loss and inductance targets. The KiCad layout is then extracted in Ansys Q3D, and the resulting network is evaluated in LTspice. A failed drain-voltage, gate-voltage, or sharing constraint returns the design to the layout or gate-drive stage. This iteration is necessary because minimizing the power loop alone does not determine the switching behavior of a paralleled stage.')
fig('workflow','Fig. 2. Iterative building-block design. Power- and gate-loop resistance R and inductance L are extracted together with mutual terms. DPT measurements are reserved for experimental correlation in the full paper.')
p('Fig. 3 shows the two board variants and the extraction geometry. Both use four copper layers, closely adjacent power planes, and four repeated cells. Distributed bypass capacitors close the fast commutation paths, while the larger bank supplies switching-frequency charge. The DPT board adds one coaxial shunt per cell. The operating board removes these shunts and provides the intended cooling arrangement.')
fig('boards','Fig. 3. Building-block implementation. (a) DPT board with individual cell-current sensing. (b) Operating board with heatsink. (c) Q3D geometry of the instrumented commutation cells; vertical dimensions are exaggerated for visibility. Board renderings denote designed hardware, not experimental operation.')
p('The cell region is extracted at 100 MHz for the power-loop inductance, and the full board at dc for the resistance matrix. Mutual inductances are retained. Capacitor equivalent series inductance and shunt inductance are added to the copper network. The effective cell loop is evaluated with all parallel cells commutating. The two boards differ primarily through the measurement path; their cell loops are symmetric at the precision supported by the extraction.')
p('The gate extraction separates the driver-to-star path from the individual gate and Kelvin branches. For simultaneous pair switching, the effective high-side and low-side gate-loop inductances are 20.6 and 12.3 nH. Their shared portions explain why branch matching can coexist with a large total gate inductance. These quantities cannot be replaced by the isolated branch inductance in a switching model.')

pg()
h('V. Switching Simulation and Design Iteration')
p('The extracted copper network is combined with eight manufacturer device models, bias-dependent bypass capacitance, component inductances, and source/sink drive networks representing the selected driver output impedances. The operating-board cases retain the preliminary 1.7 nH lumped gate model. The DPT-board cases include the extracted gate network and its component interconnects. All cases use the candidate submodule voltage in Table I; the saved DPT runs reach approximately 316 A at the first turn-off. Thus they evaluate the vicinity of the design peak without claiming an exact peak-current test.')
p('Fig. 4(a) shows the four DPT cell currents during the second turn-on. The nominal symmetric model gives equal switching-energy sharing at the reporting precision. Fig. 4(b) shows the highest stressed high-side transistor in each configuration. The critical drain-voltage event occurs when the opposite switch turns on. Fig. 4(c) compares the held-off gate response. The effective turn-on resistance is denoted Rg,on; for a DPT pair it includes the shared resistor under simultaneous switching. Turn-on energy Eon is integrated over the commutation interval and averaged over the four active devices.')
fig('switching','Fig. 4. Extracted-network DPT simulations. (a) DPT cell currents, with cells distinguished by line style. (b) Worst-cell high-side drain voltage: operating board at 2 Ω (black solid) and 3 Ω (teal dashed), DPT board at 3.5 Ω (wine dash-dot) and 2.5 Ω (ochre dotted). (c) High-side gate-to-source voltage of cell 1 for the baseline operating board and adjusted DPT board, with the same case styles. Grey dotted lines mark the drain-voltage target and minimum gate threshold. Time is relative to the second turn-on command.')
tab('III','Extracted-network simulation results',['Board and gate model','Effective loop (nH)','Rg,on (Ω)','Peak VDS (V)','Mean Eon (µJ)'],[
['Operating, lumped gate','0.41','2.0','91','27.3'],['Operating, lumped gate','0.41','3.0','86','32.8'],['DPT, extracted gate','0.65','2.5','105','24.3'],['DPT, extracted gate','0.65','3.5','93','31.2']], [6.0,2.5,2.0,2.5,3.0])
p('Table III quantifies the tradeoff. Increasing turn-on resistance reduces the operating-board drain stress at the cost of higher turn-on energy. The extracted DPT gate network produces a higher peak than the lumped-gate assumption and requires slower turn-on. Its adjusted peak remains below the device blocking rating but exceeds the design target in Table I. This result identifies remaining margin work rather than establishing compliance of the instrumented board.')
p('With zero off-state bias, the simulated held-off gate can exceed the minimum threshold specified in [9]. The DPT setting uses −1.9 V off-state bias; the gate remains below threshold, while the approximately −3.4 V negative ringing remains within the −4 V limit. Shortening the shared gate path is the direct layout remedy. Additional slowing must be assessed against switching and reverse-conduction losses.')
p('For the selected parallel count, the preliminary fundamental-period loss is dominated by conduction: approximately 12.5 W of an 18 W leg total, compared with approximately 5.0 W switching loss. This explains the preference for low on-resistance and adequate paralleling. These operating losses precede the final gate adjustment; the peak-current energies in Table III are not fundamental-period loss estimates.')

pg()
h('VI. DPT Verification in the Full Paper')
p('The DPT board provides individual cell-current sensing and local access to drain and gate voltages. The full paper will include measured turn-on and turn-off waveforms, cell-current sharing, switching energies, drain overshoot, and held-off gate voltage. Measurements will be compared with the extracted model at identical voltage, current, and gate settings. The shunts and probe connections will be retained in the test model because the measurement arrangement changes the commutation loop. Experimental verification will also determine whether the remaining DPT voltage margin requires a further layout or drive revision.')
h('VII. Conclusion')
p('The SPB building block is governed by the coupling between device voltage class, parallel count, and parasitic impedance. A charge-based device figure of merit selects EPC2361 among the surveyed devices, while loss and voltage-stress calculations motivate four transistors per switch. Symmetric cells and one gate-driver channel per transistor pair address the dominant sharing sensitivities. Extraction and switching simulation show that a low and matched commutation loop does not eliminate the influence of shared gate inductance. The operating-board power network meets the nominal drain-voltage target after slower turn-on, whereas the instrumented board still requires additional margin. The resulting DPT platform enables the experimental correlation planned for the full paper.')
h('References')
refs=[
'[1] S. Norrga, L. Jin, O. Wallmark, A. Mayer, and K. Ilves, “A novel inverter topology for compact EV and HEV drive systems,” in Proc. IECON, 2013, pp. 6590–6595.',
'[2] L. Jin, S. Norrga, O. Wallmark, and M. Nikouei Harnefors, “Control and modulation of the stacked polyphase bridges inverter,” in Proc. IEEE ECCE, 2014.',
'[3] L. Jin, S. Norrga, O. Wallmark, and N. Apostolopoulos, “Modulation and power losses of a stacked polyphase bridge converter,” IEEE J. Emerg. Sel. Topics Power Electron., vol. 5, no. 1, pp. 409–418, Mar. 2017.',
'[4] M. Nikouie, O. Wallmark, L. Jin, L. Harnefors, and H.-P. Nee, “DC-link stability analysis and controller design for the stacked polyphase bridges converter,” IEEE Trans. Power Electron., vol. 32, no. 2, pp. 1666–1674, Feb. 2017.',
'[5] G. Rohner, J. Huber, S. Mirić, and J. W. Kolar, “Comparative evaluation of three-phase three-level flying capacitor and stacked polyphase bridge GaN inverter systems for integrated motor drives,” Electronics, vol. 13, no. 7, Art. no. 1259, 2024, doi: 10.3390/electronics13071259.',
'[6] J. A. Brothers and T. Beechner, “GaN module design recommendations based on the analysis of a commercial 3-phase GaN module,” in Proc. IEEE ECCE, 2019, pp. 4109–4116.',
'[7] J. Lu and D. Chen, “Paralleling GaN E-HEMTs in 10 kW–100 kW systems,” in Proc. IEEE APEC, 2017, pp. 3049–3056.',
'[8] M. T. Tran et al., “A high-performance GaN power module with parallel packaging for high-current and low-voltage traction inverter applications,” IEEE J. Emerg. Sel. Topics Power Electron., vol. 13, no. 1, pp. 1188–1209, Feb. 2025.',
'[9] Efficient Power Conversion, “EPC2361 eGaN FET datasheet,” Feb. 20, 2026. [Online]. Available: https://epc-co.com/epc/Portals/0/epc/documents/datasheets/EPC2361_datasheet.pdf',
'[10] Infineon Technologies, “2EDF7275K data sheet,” rev. 2.8. [Online]. Available: https://www.infineon.com/part/2EDF7275K']
for r in refs:
 pa=p(r);pa.paragraph_format.first_line_indent=Cm(-.55);pa.paragraph_format.left_indent=Cm(.55);pa.paragraph_format.space_after=Pt(4)
footer=sec.footer.paragraphs[0];footer.alignment=1
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=title.text;doc.core_properties.subject='Anonymous ICPE conference digest'
for cp in ['comments','keywords']:setattr(doc.core_properties,cp,'')
# Let Word paginate continuously; retain the venue minimum type size and all content.
for br in doc.element.xpath('.//w:br[@w:type="page"]'):
 br.getparent().remove(br)
for pa in doc.paragraphs:
 if not pa.text and not pa._p.xpath('.//w:drawing | .//m:oMath'):
  pa.paragraph_format.space_after=Pt(0);pa.paragraph_format.space_before=Pt(0)
  pa.paragraph_format.line_spacing=Pt(1)
  for rr in pa.runs:rr.font.size=Pt(1)
for pa in doc.paragraphs:
 for rr in pa.runs:
  if 'Prior GaN studies' in rr.text:rr.text=rr.text.replace('Prior GaN studies','Prior gallium nitride (GaN) studies')
# Consistent engineering symbols in body, captions, and table headers.
from copy import deepcopy
symbols={'VDS,pk':('V','DS,pk'),'RDS(on)':('R','DS(on)'),'Coss,tr':('C','oss,tr'),'Iph,rms':('I','ph,rms'),'Vdc':('V','dc'),'Vm':('V','m'),'fsw':('f','sw'),'Qoss':('Q','oss'),'Vtest':('V','test'),'Lloop':('L','loop'),'Rg,on':('R','g,on'),'VDS':('V','DS'),'Eon':('E','on'),'Pmin':('P','min')}
pat=re.compile('|'.join(re.escape(k) for k in sorted(symbols,key=len,reverse=True)))
paras=list(doc.paragraphs)+[pp for t in doc.tables for row in t.rows for cc in row.cells for pp in cc.paragraphs]
for pa in paras:
 for rr in list(pa.runs):
  old=rr.text
  if not pat.search(old):continue
  base=deepcopy(rr._r.rPr) if rr._r.rPr is not None else None
  pos=0
  for m in pat.finditer(old):
   parts=[(old[pos:m.start()],False,False),(symbols[m.group()][0],False,True),(symbols[m.group()][1],True,False)]
   for txt,isub,ital in parts:
    if not txt:continue
    nr=OxmlElement('w:r')
    if base is not None:nr.append(deepcopy(base))
    rp=nr.find(qn('w:rPr'))
    if rp is None:rp=OxmlElement('w:rPr');nr.insert(0,rp)
    if isub:
     v=OxmlElement('w:vertAlign');v.set(qn('w:val'),'subscript');rp.append(v)
    if ital:rp.append(OxmlElement('w:i'))
    tt=OxmlElement('w:t');tt.set(qn('xml:space'),'preserve');tt.text=txt;nr.append(tt);rr._r.addprevious(nr)
   pos=m.end()
  if pos<len(old):
   nr=deepcopy(rr._r)
   for tt in list(nr):
    if tt.tag!=qn('w:rPr'):nr.remove(tt)
   tt=OxmlElement('w:t');tt.set(qn('xml:space'),'preserve');tt.text=old[pos:];nr.append(tt);rr._r.addprevious(nr)
  rr._r.getparent().remove(rr._r)
# Keep the short sensitivity table intact and repeat table headers if needed.
for table in doc.tables:
 rep=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(rep)
for row in doc.tables[1].rows[:-1]:
 for cell in row.cells:
  for pp in cell.paragraphs:pp.paragraph_format.keep_with_next=True
doc.save(OUT)
print(OUT)
print('Words:',sum(len(x.text.split()) for x in doc.paragraphs))

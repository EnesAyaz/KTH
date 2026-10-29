from pathlib import Path
import subprocess
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Cm,Pt
P=Path(r'C:\Github\KTH\PCB-Design\reports\paper');A=P/'icpe'
pdfmetrics.registerFont(TTFont('QTimes',r'C:\Windows\Fonts\times.ttf'))
c=canvas.Canvas(str(A/'q3d_geometry_views.pdf'),pagesize=(454,295))
def txt(x,y,t,size=10):c.setFont('QTimes',size);c.setFillColor('#111111');c.drawString(x,y,t)
for x,y,w,h,name,label in [(0,159,221,123,'q3d_loop_load.png','(a) Operating-board power network'),(233,159,221,123,'q3d_loop_dpt.png','(b) DPT-board power network'),(0,8,257,140,'q3d_model_cells.png','(c) Detail of one cell pair')]:
 c.drawImage(str(P/'figures/revb'/name),x,y,width=w,height=h,preserveAspectRatio=True,anchor='c',mask='auto');txt(x+5,y+h+3,label,9)
for yy,label,col in [(115,'DC positive','#d92c27'),(96,'DC negative','#2167b7'),(77,'Switch node','#ed9c20'),(58,'Source-to-shunt path','#329f50'),(39,'Device / capacitor terminals','#d5b85b')]:
 c.setFillColor(col);c.rect(277,yy,8,7,fill=1,stroke=0);txt(291,yy,label,9)
c.save()
pop=r'C:\Users\enesa\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe'
subprocess.run([pop,'-singlefile','-r','450','-png',str(A/'q3d_geometry_views.pdf'),str(A/'q3d_geometry_views')],check=True)
subprocess.run([r'C:\Program Files\Inkscape\bin\inkscape.com',str(A/'q3d_geometry_views.pdf'),'--export-type=svg','--export-plain-svg','--export-filename='+str(A/'q3d_geometry_views.svg')],check=True)
doc=Document(P/'SPB_ICPE_complete_manuscript_2026-10-29.docx')
# Renumber existing switching figure and its references while retaining all run formatting.
for q in doc.paragraphs:
 for r in q.runs:
  if 'Fig. 8' in r.text:r.text=r.text.replace('Fig. 8','Fig. 9')
anchor=next(q for q in doc.paragraphs if q.text=='C. Switching Assessment and Design Feedback')
q=doc.add_paragraph('Fig. 8 shows the Q3D conductor models underlying the extraction. The operating-board view is compared with the DPT-board network, which includes the additional low-side source-to-shunt path. The cell-pair detail shows the overlapping conductors, interlayer connections, and device and capacitor terminals used to define the local power network. These are geometry views rather than field-distribution plots.');q.alignment=3;anchor._p.addprevious(q._p)
q=doc.add_paragraph();q.alignment=1;q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.space_after=Pt(0);q.paragraph_format.keep_with_next=True;q.add_run().add_picture(str(A/'q3d_geometry_views.png'),width=Cm(16));anchor._p.addprevious(q._p)
q=doc.add_paragraph('Fig. 8. Ansys Q3D conductor models used for parasitic extraction: (a) operating-board power network; (b) instrumented DPT-board power network, including the additional source-to-shunt conductors; and (c) detail of one cell pair showing the local conductor stack, interlayer connections, and device and capacitor terminals. Colors identify electrical nets and terminals, not field magnitude.','Caption');q.paragraph_format.keep_together=True;anchor._p.addprevious(q._p)
out=P/'SPB_ICPE_complete_manuscript_with_Q3D_2026-10-29.docx';doc.save(out);print(out)
